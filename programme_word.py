"""Exports Word non destructifs des projets et programmes bilingues."""
from io import BytesIO
from docx import Document


def programme_docx_bytes(rows, title="Programme liturgique", provisional=True):
    """Construit un document à partir des lignes, sans mutation de l'état."""
    doc = Document()
    doc.add_heading(title, level=0)
    if provisional:
        doc.add_paragraph("PROJET — NON VALIDÉ / NE PAS DIFFUSER", style="Subtitle")
    for row in rows:
        doc.add_heading(
            f"{row.get('celebration', 'Célébration')} — {row.get('Dimanche', row.get('date', ''))}",
            level=1,
        )
        for label in ("Références", "Lecteurs", "Monition + P.U."):
            value = str(row.get(label, "")).strip()
            if value:
                doc.add_paragraph(label, style="Heading 2")
                doc.add_paragraph(value)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
