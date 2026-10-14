import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


def build_word_document(meta, monition, pu_intro, intentions, pu_conclusion, response):
    """Construit le document final avec uniquement la monition et les 4 intentions."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    date_label = str(meta.get("date_label", "")).strip()

    celebration = str(meta.get("celebration") or "").strip()
    references = str(meta.get("references") or "").strip()
    if celebration:
        top = document.add_paragraph()
        top.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = top.add_run(celebration.upper())
        run.bold = True
        run.font.size = Pt(12)
    if date_label or references:
        details = document.add_paragraph()
        details.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = details.add_run("  •  ".join(filter(None, [date_label, references])))
        run.italic = True
        run.font.size = Pt(9)
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("MONITION")
    run.bold = True
    run.font.size = Pt(14)

    p = document.add_paragraph(str(monition or "").strip())
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("PRIÈRE UNIVERSELLE")
    run.bold = True
    run.font.size = Pt(14)

    for idx, intention in enumerate(list(intentions or [])[:4], start=1):
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        lead = p.add_run(f"{idx}. ")
        lead.bold = True
        content = str(intention or "").strip()
        ending = "Prions le Seigneur."
        if content.endswith(ending):
            p.add_run(content[:-len(ending)].rstrip() + " ")
            p.add_run(ending).bold = True
        else:
            p.add_run(content)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("Programme liturgique — document de préparation")
    fr.font.size = Pt(8)

    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    return output.getvalue()
