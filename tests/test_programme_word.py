import io
import unittest
from copy import deepcopy
from docx import Document
from programme_word import programme_docx_bytes


class ProgrammeWordTests(unittest.TestCase):
    def test_provisional_word_contains_programme_but_not_announcements(self):
        rows = [{
            "date": "2026-10-31", "Dimanche": "31/10/2026",
            "celebration": "Samedi soir", "Références": "1re : R1",
            "Lecteurs": "FR : A\nMO : B",
            "Monition + P.U.": "FR : C\nMO : D",
            "codes": {"r1": "F1", "r2": "M1", "f_mon": "F2", "m_mon": "M2"},
        }]
        original = deepcopy(rows)
        data = programme_docx_bytes(rows, "Novembre 2026")
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("PROJET — NON VALIDÉ", text)
        self.assertIn("Samedi soir", text)
        self.assertIn("Monition + P.U.", text)
        self.assertNotIn("Annonces", text)
        self.assertEqual(rows, original)

    def test_published_word_has_no_provisional_banner(self):
        data = programme_docx_bytes([], provisional=False)
        doc = Document(io.BytesIO(data))
        self.assertNotIn("NON VALIDÉ", "\n".join(p.text for p in doc.paragraphs))


if __name__ == "__main__":
    unittest.main()
