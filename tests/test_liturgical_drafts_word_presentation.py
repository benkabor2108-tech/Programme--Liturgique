import io
import unittest
from docx import Document

from liturgical_drafts_word import build_word_document


class ReaderDocumentPresentationTests(unittest.TestCase):
    def test_monition_paragraphs_and_empty_intentions(self):
        doc = Document(io.BytesIO(build_word_document(
            {"celebration": "Célébration de test", "date_label": "18/10/2026"},
            "Premier paragraphe.\n\nDeuxième paragraphe.", "",
            ["Pour l'Église. Prions le Seigneur.", "", "Pour les malades. Prions le Seigneur."],
            "", "",
        )))
        paragraphs = [p.text for p in doc.paragraphs]
        self.assertIn("Premier paragraphe.", paragraphs)
        self.assertIn("Deuxième paragraphe.", paragraphs)
        self.assertFalse(any(p.strip() == "2." for p in paragraphs))
        self.assertTrue(any(p.startswith("2. Pour les malades.") for p in paragraphs))
        self.assertFalse(any("document de préparation" in p.text for p in doc.sections[0].footer.paragraphs))

    def test_word_adds_missing_prayer_endings_in_bold(self):
        raw = build_word_document(
            {}, "Monition.", "",
            ["Pour les familles.", "Pour la paix. Prions le Seigneur.",
             "Pour les malades.", "Pour les jeunes."], "", "",
        )
        doc = Document(io.BytesIO(raw))
        prayer_paragraphs = [p for p in doc.paragraphs if "Prions le Seigneur." in p.text]
        self.assertEqual(len(prayer_paragraphs), 4)
        for paragraph in prayer_paragraphs:
            self.assertEqual(paragraph.text.count("Prions le Seigneur."), 1)
            self.assertTrue(any(run.bold and run.text == "Prions le Seigneur." for run in paragraph.runs))

    def test_reader_document_matches_layout_and_excludes_celebrant(self):
        meta = {
            "celebration": "29e dimanche du Temps ordinaire — Année A",
            "date_label": "18/10/2026",
            "refs": {"r1": "Is 45, 1.4-6", "ps": "Ps 95", "r2": "1 Th 1, 1-5b", "ev": "Mt 22, 15-21"},
        }
        prayers = [f"Pour notre communauté, intention {n}. Prions le Seigneur." for n in range(1, 5)]
        raw = build_word_document(meta, "Frères et sœurs, accueillons la Parole.", "INTRODUCTION DU CÉLÉBRANT", prayers, "CONCLUSION DU CÉLÉBRANT", "Réponse")
        doc = Document(io.BytesIO(raw))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("29E DIMANCHE", text)
        self.assertIn("Mt 22, 15-21", text)
        self.assertIn("MONITION", text)
        self.assertIn("PRIÈRE UNIVERSELLE", text)
        self.assertNotIn("INTRODUCTION DU CÉLÉBRANT", text)
        self.assertNotIn("CONCLUSION DU CÉLÉBRANT", text)
        self.assertEqual(text.count("Prions le Seigneur."), 4)
        for p in doc.paragraphs:
            if "Prions le Seigneur." in p.text:
                self.assertTrue(any(r.bold and "Prions le Seigneur." in r.text for r in p.runs))


if __name__ == "__main__":
    unittest.main()
